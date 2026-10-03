# -*- coding: utf-8 -*-
"""
Script to generate the comprehensive Word Document for TV1:
BÁO CÁO TOÀN DIỆN VÀ CHI TIẾT TỪ A-Z: NỀN TẢNG VÀ BẢN ĐỒ TƯ DUY TÍNH TOÁN BẤT ĐỒNG BỘ
Học phần: Tính toán song song và phân tán (Parallel and Distributed Computing)
Thành viên: TV1 (Workstream WS-01) - Chủ trì Nền tảng & Bản đồ Khái niệm
"""

import sys
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D5D8DC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    run.font.name = "Arial"
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(16)
        run.bold = True
        run.font.color.rgb = RGBColor(192, 57, 43) # Coral Red / Crimson
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(13)
        run.bold = True
        run.font.color.rgb = RGBColor(44, 62, 80) # Slate Navy
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(11.5)
        run.bold = True
        run.font.color.rgb = RGBColor(52, 73, 94)
    return h

def add_para(doc, text, bold_prefix=None, italic_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(33, 47, 61)
    if italic_prefix:
        r_it = p.add_run(italic_prefix)
        r_it.italic = True
        r_it.font.name = "Arial"
        r_it.font.size = Pt(10)
        r_it.font.color.rgb = RGBColor(127, 140, 141)
    r_body = p.add_run(text)
    r_body.font.name = "Arial"
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = RGBColor(44, 62, 80)
    return p

def add_bullet(doc, text, bold_prefix=None, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(33, 47, 61)
    r_body = p.add_run(text)
    r_body.font.name = "Arial"
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = RGBColor(44, 62, 80)
    return p

def add_callout(doc, text, title="GHI CHÚ CHUYÊN SÂU", bg_hex="FDEDEC", border_hex="E74C3C"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"📌 {title}: ")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(192, 57, 43)
    
    run_b = p.add_run(text)
    run_b.font.name = "Arial"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(44, 62, 80)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_code_block(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        run_cap = p_cap.add_run(f"💻 Đoạn mã minh họa: {caption}")
        run_cap.italic = True
        run_cap.font.name = "Arial"
        run_cap.font.size = Pt(9)
        run_cap.font.color.rgb = RGBColor(127, 140, 141)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F4F6F6")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="BDC3C7"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="7F8C8D"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="BDC3C7"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="BDC3C7"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(33, 47, 61)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def create_styled_table(doc, headers, rows_data, col_widths=None, header_bg="E74C3C"):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.runs[0]
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "FFFFFF" if r_idx % 2 == 0 else "FDF2E9" # subtle tint
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            if len(p.runs) > 0:
                p.runs[0].font.name = "Arial"
                p.runs[0].font.size = Pt(9)
                p.runs[0].font.color.rgb = RGBColor(44, 62, 80)
                
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table

print("Document formatting primitives loaded.")
