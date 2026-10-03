# -*- coding: utf-8 -*-
"""
build_tv1_clean_original_presentation.py
Tạo bộ slide PowerPoint CHUẨN XÁC 100% THEO FILE BÀI GIẢNG GỐC (BaiGiang_TinhToanSongSong_PhanTan_C4.pdf)
- Đúng phạm vi của TV1: PHẦN 1: GIỚI THIỆU (Từ Slide 1 đến Slide 8 của file gốc).
- Loại bỏ hoàn toàn nội dung ngoài lề (không Amdahl toán học phức tạp, không checklist benchmark, không quy trình GitHub).
- Giữ đúng cấu trúc slide gốc: Tuần tự, Đồng thời (PID, xen kẽ), Song song (đa nhân, độc lập), Bất đồng bộ (không theo thứ tự, ví dụ Đi thi / Bãi giữ xe, mô hình luồng đơn thực thi đồng thời tạm dừng và tiếp tục lại).
- Tỷ lệ 4:3 chuẩn mẫu C4 (793.8 pt x 595.2 pt), Banner đỏ #E74C3C, Footer 3 khối chuẩn.
- Nhúng 6 sơ đồ minh họa trực quan chất lượng cao bám sát nội dung bài giảng gốc.
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

COLOR_RED = RGBColor(231, 76, 60)       # #E74C3C
COLOR_GRAY = RGBColor(189, 195, 199)    # #BDC3C7
COLOR_DARK = RGBColor(44, 62, 80)       # #2C3E50
COLOR_BODY = RGBColor(33, 47, 61)       # #212F3D
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_CARD_BG = RGBColor(248, 249, 250) # #F8F9FA
COLOR_GREEN = RGBColor(39, 174, 96)     # #27AE60
COLOR_ORANGE = RGBColor(230, 126, 34)   # #E67E22
COLOR_BLUE = RGBColor(41, 128, 185)     # #2980B9

IMG_DIR = "slide_images"

def create_slide_base(prs, slide_num, header_title, footer_subtitle, session_info="01 - 2026"):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 1. Header đỏ phía trên
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(0), Pt(14.16), Pt(765.36), Pt(99.24))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_RED
    top_bar.line.fill.background()
    
    tf_top = top_bar.text_frame
    tf_top.word_wrap = True
    tf_top.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_top.margin_left = Pt(28.34)
    p_top = tf_top.paragraphs[0]
    p_top.text = header_title.upper()
    p_top.font.name = "Arial"
    p_top.font.size = Pt(26)
    p_top.font.bold = True
    p_top.font.color.rgb = COLOR_WHITE

    # 2. Footer trái (Số trang)
    left_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(14.16), Pt(538.56), Pt(42.48), Pt(42.6))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = COLOR_RED
    left_box.line.fill.background()
    tf_num = left_box.text_frame
    tf_num.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_num = tf_num.paragraphs[0]
    p_num.text = str(slide_num)
    p_num.font.name = "Arial"
    p_num.font.size = Pt(16)
    p_num.font.bold = True
    p_num.font.color.rgb = COLOR_WHITE
    p_num.alignment = PP_ALIGN.CENTER

    # 3. Footer giữa (Phụ đề chuyên đề)
    center_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(70.92), Pt(538.56), Pt(510.24), Pt(42.6))
    center_box.fill.solid()
    center_box.fill.fore_color.rgb = COLOR_GRAY
    center_box.line.fill.background()
    tf_mid = center_box.text_frame
    tf_mid.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_mid.margin_left = Pt(15.0)
    p_mid = tf_mid.paragraphs[0]
    p_mid.text = footer_subtitle
    p_mid.font.name = "Arial"
    p_mid.font.size = Pt(12.5)
    p_mid.font.color.rgb = COLOR_DARK

    # 4. Footer phải (Học kỳ / Niên khóa)
    right_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(595.32), Pt(538.56), Pt(198.36), Pt(42.6))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = COLOR_RED
    right_box.line.fill.background()
    tf_right = right_box.text_frame
    tf_right.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_right = tf_right.paragraphs[0]
    p_right.text = session_info
    p_right.font.name = "Arial"
    p_right.font.size = Pt(13)
    p_right.font.bold = True
    p_right.font.color.rgb = COLOR_WHITE
    p_right.alignment = PP_ALIGN.CENTER

    return slide

def add_content_title(slide, text):
    tx_box = slide.shapes.add_textbox(Pt(28.34), Pt(122.0), Pt(735.0), Pt(32.0))
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_top = Pt(0)
    tf.margin_left = Pt(0)
    p = tf.paragraphs[0]
    p.text = f"➢ {text}"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK
    return tx_box

def add_bullets(slide, items, left=Pt(28.34), top=Pt(155.0), width=Pt(735.0), height=Pt(160.0), font_size=Pt(12.5)):
    tx_box = slide.shapes.add_textbox(left, top, width, height)
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_top = Pt(0)
    tf.margin_left = Pt(0)
    
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        p.space_before = Pt(2)
        
        if isinstance(item, tuple):
            prefix, body = item
            r_pre = p.add_run()
            r_pre.text = f"□ {prefix}: "
            r_pre.font.name = "Arial"
            r_pre.font.size = font_size
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK
            
            r_body = p.add_run()
            r_body.text = body
            r_body.font.name = "Arial"
            r_body.font.size = font_size
            r_body.font.color.rgb = COLOR_BODY
        else:
            r = p.add_run()
            r.text = f"□ {item}"
            r.font.name = "Arial"
            r.font.size = font_size
            r.font.color.rgb = COLOR_BODY
    return tx_box

def add_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = f"[GIÁO ÁN GIẢNG DẠY]:\n{notes_text}"

def add_table_box(slide, headers, rows, left, top, width, height, col_widths=None):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    table = table_shape.table
    
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_RED
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_value in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            bg_col = COLOR_WHITE if row_idx % 2 == 0 else COLOR_CARD_BG
            cell.fill.fore_color.rgb = bg_col
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = str(cell_value)
            p.font.name = "Arial"
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_BODY
            
    if col_widths:
        for idx, w in enumerate(col_widths):
            table.columns[idx].width = Pt(w)
            
    return table_shape

def build_tv1_clean_presentation(prs):
    # =========================================================================
    # SLIDE 1: BÌA CHƯƠNG 4 (Khớp chính xác Slide 1 file gốc)
    # =========================================================================
    blank_layout = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank_layout)
    
    tx_top = s1.shapes.add_textbox(Pt(28.34), Pt(45.0), Pt(735.0), Pt(65.0))
    tf_top = tx_top.text_frame
    p_t1 = tf_top.paragraphs[0]
    p_t1.text = "TRƯỜNG ĐẠI HỌC KỸ THUẬT - CÔNG NGHỆ CẦN THƠ (CTUET)"
    p_t1.font.name = "Arial"
    p_t1.font.size = Pt(14)
    p_t1.font.bold = True
    p_t1.font.color.rgb = COLOR_DARK
    
    p_t2 = tf_top.add_paragraph()
    p_t2.text = "KHOA CÔNG NGHỆ THÔNG TIN — HỌC PHẦN: TÍNH TOÁN SONG SONG VÀ PHÂN TÁN"
    p_t2.font.name = "Arial"
    p_t2.font.size = Pt(11.5)
    p_t2.font.color.rgb = COLOR_GRAY

    # Băng đỏ giữa trang
    mid_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(0), Pt(235.0), Pt(765.36), Pt(125.0))
    mid_bar.fill.solid()
    mid_bar.fill.fore_color.rgb = COLOR_RED
    mid_bar.line.fill.background()
    
    tf_mid = mid_bar.text_frame
    tf_mid.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_mid.margin_left = Pt(28.34)
    p_m1 = tf_mid.paragraphs[0]
    p_m1.text = "CHƯƠNG  4:"
    p_m1.font.name = "Arial"
    p_m1.font.size = Pt(26)
    p_m1.font.bold = True
    p_m1.font.color.rgb = COLOR_WHITE
    
    p_m2 = tf_mid.add_paragraph()
    p_m2.text = "TÍNH TOÁN BẤT ĐỒNG BỘ TRONG PYTHON"
    p_m2.font.name = "Arial"
    p_m2.font.size = Pt(30)
    p_m2.font.bold = True
    p_m2.font.color.rgb = COLOR_WHITE

    tx_bot = s1.shapes.add_textbox(Pt(28.34), Pt(385.0), Pt(735.0), Pt(140.0))
    tf_bot = tx_bot.text_frame
    p_b1 = tf_bot.paragraphs[0]
    p_b1.text = "PHẦN I: GIỚI THIỆU (NỀN TẢNG CÁC MÔ HÌNH THỰC THI)"
    p_b1.font.name = "Arial"
    p_b1.font.size = Pt(14)
    p_b1.font.bold = True
    p_b1.font.color.rgb = COLOR_DARK
    
    p_b2 = tf_bot.add_paragraph()
    p_b2.text = "Người thực hiện: Thành viên 1 (Phụ trách Slide 1 đến Slide 8 bài giảng gốc)"
    p_b2.font.name = "Arial"
    p_b2.font.size = Pt(12.5)
    p_b2.font.bold = True
    p_b2.font.color.rgb = COLOR_ORANGE
    
    p_b3 = tf_bot.add_paragraph()
    p_b3.text = "Giảng viên hướng dẫn: ThS. Lê Anh Nhã Uyên (lanuyen@ctuet.edu.vn)"
    p_b3.font.name = "Arial"
    p_b3.font.size = Pt(11)
    p_b3.font.color.rgb = COLOR_BODY

    add_speaker_notes(s1, "Kính chào cô và các bạn. Hôm nay em đại diện cho TV1 trình bày Phần I: Giới thiệu của Chương 4 - Tính toán bất đồng bộ trong Python, bao gồm 4 mô hình nền tảng: Tuần tự, Đồng thời, Song song và Bất đồng bộ.")

    # =========================================================================
    # SLIDE 2: TỔNG QUAN NỘI DUNG PHẦN GIỚI THIỆU
    # =========================================================================
    s2 = create_slide_base(prs, 2, "GIỚI THIỆU", "Tổng quan 4 Mô hình Thực thi Nền tảng")
    add_content_title(s2, "Bốn mô hình tính toán nền tảng của Chương 4")
    add_bullets(s2, [
        ("Mô hình 1 — Xử lý Tuần tự (Sequential)", "Xử lý 1 tác vụ trong 1 khoảng thời gian; các task nối tiếp nhau theo thứ tự cố định."),
        ("Mô hình 2 — Xử lý Đồng thời (Concurrency)", "Phân chia và điều phối nhiều tác vụ cùng tiến triển xen kẽ nhau trên 1 nhân CPU."),
        ("Mô hình 3 — Xử lý Song song (Parallelism)", "Nhiều tác vụ độc lập được thực thi đồng thời vật lý tại cùng thời điểm trên nhiều nhân CPU."),
        ("Mô hình 4 — Xử lý Bất đồng bộ (Asynchronous)", "Xử lý không theo thứ tự, giải quyết nhiều yêu cầu đồng thời trong khoảng thời gian ngắn hơn.")
    ], top=Pt(150.0), height=Pt(170.0), font_size=Pt(12.5))
    s2.shapes.add_picture(os.path.join(IMG_DIR, "diag_orig_compare_3_models.png"), Pt(28.34), Pt(325.0), Pt(735.0), Pt(195.0))

    add_speaker_notes(s2, "Trước khi bước vào các thư viện cụ thể như concurrent.futures hay asyncio, chúng ta cần nắm vững 4 mô hình thực thi nền tảng này để hiểu rõ bản chất máy tính xử lý công việc ra sao.")

    # =========================================================================
    # SLIDE 3: XỬ LÝ TUẦN TỰ (SEQUENTIAL) — Trang 2 slide gốc
    # =========================================================================
    s3 = create_slide_base(prs, 3, "GIỚI THIỆU", "Sequential - Tuần tự (Slide 2 gốc)")
    add_content_title(s3, "Xử lý tuần tự (Sequential Computing)")
    add_bullets(s3, [
        ("Định nghĩa xử lý tuần tự", "Xử lý 1 tác vụ (task) trong một khoảng thời gian xác định."),
        ("Thứ tự thực thi", "Các task sẽ được thực thi tuần tự theo thứ tự: Tác vụ trước phải hoàn thành xong thì tác vụ sau mới được bắt đầu."),
        ("Đặc điểm vận hành", "Chương trình chỉ duy trì một dòng thực thi duy nhất; không có sự chia sẻ hay xen kẽ giữa các tác vụ."),
        ("Hạn chế chính", "Nếu một tác vụ bị nghẽn (ví dụ chờ dữ liệu nhập xuất I/O), toàn bộ các tác vụ phía sau đều bị chặn đứng và phải chờ đợi.")
    ], top=Pt(150.0), height=Pt(165.0), font_size=Pt(12.5))
    s3.shapes.add_picture(os.path.join(IMG_DIR, "diag_orig_sequential.png"), Pt(28.34), Pt(330.0), Pt(735.0), Pt(190.0))

    add_speaker_notes(s3, "Đây là nội dung Slide 2 bài giảng gốc: Xử lý tuần tự là mô hình cơ bản nhất. 1 task làm xong mới tới task tiếp theo. Nhược điểm là nếu task 1 mất thời gian chờ đĩa hoặc mạng, CPU sẽ bị nhàn rỗi lãng phí.")

    # =========================================================================
    # SLIDE 4: XỬ LÝ ĐỒNG THỜI (CONCURRENCY) — VAI TRÒ HĐH & TIẾN TRÌNH (Trang 3 slide gốc)
    # =========================================================================
    s4 = create_slide_base(prs, 4, "GIỚI THIỆU", "Concurrency - Đồng thời (Slide 3 gốc)")
    add_content_title(s4, "Xử lý đồng thời: Vai trò của Hệ điều hành và Tiến trình")
    add_bullets(s4, [
        ("Hệ điều hành quản lý chương trình", "Mỗi chương trình chạy trong hệ thống tương ứng với một tiến trình (Process)."),
        ("Định danh tiến trình (PID)", "Hệ điều hành cấp phát một mã số Process ID (PID) duy nhất để nhận diện, theo dõi và quản lý tài nguyên."),
        ("Cơ chế phân chia của CPU Core", "Nhân CPU (CPU Core) chia các task của process thành các task nhỏ hơn + sắp xếp xen kẽ nhau."),
        ("Mục tiêu tối ưu", "Tận dụng tối đa thời gian rảnh của task này để chuyển sang thực hiện task khác, không để CPU rơi vào trạng thái lãng phí.")
    ], top=Pt(150.0), height=Pt(165.0), font_size=Pt(12.5))
    s4.shapes.add_picture(os.path.join(IMG_DIR, "diag_orig_concurrency_pid.png"), Pt(28.34), Pt(325.0), Pt(735.0), Pt(195.0))

    add_speaker_notes(s4, "Nội dung Slide 3 bài giảng gốc: HĐH quản lý các chương trình thông qua Process ID (PID). CPU Core thông minh ở chỗ nó băm nhỏ task và sắp xếp xen kẽ nhau để tận dụng thời gian rảnh.")

    # =========================================================================
    # SLIDE 5: XỬ LÝ ĐỒNG THỜI — NGUYÊN LÝ THỰC THI TRÊN 1 NHÂN CPU (Trang 4 slide gốc)
    # =========================================================================
    s5 = create_slide_base(prs, 5, "GIỚI THIỆU", "Concurrency - Đồng thời (Slide 4 gốc)")
    add_content_title(s5, "Xử lý đồng thời: Điều phối nhiều tác vụ trong cùng khoảng thời gian")
    add_bullets(s5, [
        ("Phân chia và điều phối tác vụ", "Hệ thống phân chia, điều phối nhiều tác vụ (task) khác nhau trong cùng một khoảng thời gian."),
        ("Quy tắc vật lý vi mô bên dưới", "Bên dưới CPU core chỉ có thể thực thi MỘT TASK NHỎ trong task lớn tại một thời điểm vi mô."),
        ("Bản chất thực tế", "Máy tính xử lý nhiều việc cùng lúc tại cùng thời điểm dưới góc nhìn của người dùng; nhưng tại 1 thời điểm vi mô trên 1 core, chỉ xử lý 1 task."),
        ("Kết luận cốt lõi", "Concurrency là sự TIẾN TRIỂN XEN KẼ (Interleaving), giải quyết bài toán quản lý nhiều việc cùng lúc trên tài nguyên hạn chế.")
    ], top=Pt(150.0), height=Pt(165.0), font_size=Pt(12.5))

    card_s5 = [
        ("Nguyên lý Concurrency", "Phân chia và sắp xếp xen kẽ nhiều tác vụ trong cùng một khoảng thời gian."),
        ("Góc nhìn người dùng (Vĩ mô)", "Cảm giác các chương trình chạy đồng thời mượt mà cùng một lúc."),
        ("Góc nhìn CPU Core (Vi mô)", "Tại mỗi thời điểm cực vi mô, 1 nhân CPU chỉ xử lý lệnh của đúng 1 task!")
    ]
    # Dùng card đẹp thay thế
    c_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(28.34), Pt(335.0), Pt(735.0), Pt(180.0))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = COLOR_CARD_BG
    c_box.line.color.rgb = COLOR_ORANGE
    c_box.line.width = Pt(1.5)
    tf_c = c_box.text_frame
    tf_c.margin_left = Pt(20); tf_c.margin_top = Pt(15)
    p0 = tf_c.paragraphs[0]
    p0.text = "TỔNG KẾT BẢN CHẤT XỬ LÝ ĐỒNG THỜI (CONCURRENCY TRÊN 1 CORE):"
    p0.font.name = "Arial"; p0.font.size = Pt(13); p0.font.bold = True; p0.font.color.rgb = COLOR_DARK
    for k, v in card_s5:
        p = tf_c.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run(); r1.text = f"• {k}: "; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = COLOR_ORANGE
        r2 = p.add_run(); r2.text = v; r2.font.size = Pt(11); r2.font.color.rgb = COLOR_BODY

    add_speaker_notes(s5, "Slide 4 bài giảng gốc chốt lại một nguyên lý cực kỳ quan trọng: Concurrency là điều phối nhiều task trong cùng 1 khoảng thời gian, nhưng tại 1 thời điểm trên 1 nhân CPU thì chỉ có 1 task được thi hành.")

    # =========================================================================
    # SLIDE 6: XỬ LÝ SONG SONG (PARALLEL) — THỰC THI ĐỒNG THỜI ĐA NHÂN (Trang 5 slide gốc)
    # =========================================================================
    s6 = create_slide_base(prs, 6, "GIỚI THIỆU", "Parallel - Song song (Slide 5 gốc)")
    add_content_title(s6, "Xử lý song song (Parallel Computing)")
    add_bullets(s6, [
        ("Định nghĩa xử lý song song", "Nhiều task khác nhau được xử lý trong CÙNG 1 THỜI ĐIỂM (thực thi đồng thời vật lý)."),
        ("Tính độc lập dữ liệu", "Các task phải hoàn toàn độc lập với nhau để có thể chạy song song mà không tranh chấp bộ nhớ."),
        ("Điều kiện phần cứng bắt buộc", "Chỉ có thể thực hiện trên máy tính có số nhân (core) vật lý lớn hơn 1 (Multi-core CPU, Multi-CPU hoặc GPU)."),
        ("Hiệu quả đạt được", "Rút ngắn thời gian thực thi thực tế của toàn bộ chương trình nhờ tận dụng sức mạnh của nhiều bộ xử lý cùng lúc.")
    ], top=Pt(150.0), height=Pt(165.0), font_size=Pt(12.5))
    s6.shapes.add_picture(os.path.join(IMG_DIR, "diag_orig_parallel_multicore.png"), Pt(28.34), Pt(325.0), Pt(735.0), Pt(195.0))

    add_speaker_notes(s6, "Slide 5 bài giảng gốc: Xử lý song song là nhiều task được xử lý thực sự cùng 1 thời điểm. Điều kiện tiên quyết: Máy tính bắt buộc phải có số core lớn hơn 1 và các task phải độc lập nhau.")

    # =========================================================================
    # SLIDE 7: XỬ LÝ SONG SONG — QUÁ TRÌNH ĐỒNG THỜI TRÊN TỪNG NHÂN (Trang 6 slide gốc)
    # =========================================================================
    s7 = create_slide_base(prs, 7, "GIỚI THIỆU", "Parallel - Song song (Slide 6 gốc)")
    add_content_title(s7, "Mối quan hệ thực tế giữa Song song và Đồng thời")
    add_bullets(s7, [
        ("Thực tế kiến trúc đa nhân", "Trên thực tế, trên mỗi nhân của CPU vẫn xảy ra quá trình xử lý ĐỒNG THỜI (xen kẽ các task nền của hệ điều hành)."),
        ("Điều kiện phân định rõ ràng", "Tại một thời điểm, KHÔNG xảy ra việc xử lý cùng một task trên hai nhân CPU khác nhau."),
        ("Sự phối hợp giữa hai mô hình", "Hệ thống đa nhân vừa có tính song song (giữa các nhân với nhau) vừa có tính đồng thời (bên trong từng nhân riêng lẻ)."),
        ("Ý nghĩa kỹ thuật", "Giúp tối ưu hóa toàn diện: Vừa khai thác hết năng lực đa lõi, vừa không lãng phí thời gian nhàn rỗi của từng lõi.")
    ], top=Pt(150.0), height=Pt(165.0), font_size=Pt(12.5))

    c_box7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(28.34), Pt(335.0), Pt(735.0), Pt(180.0))
    c_box7.fill.solid()
    c_box7.fill.fore_color.rgb = COLOR_CARD_BG
    c_box7.line.color.rgb = COLOR_GREEN
    c_box7.line.width = Pt(1.5)
    tf_7 = c_box7.text_frame
    tf_7.margin_left = Pt(20); tf_7.margin_top = Pt(15)
    p7 = tf_7.paragraphs[0]
    p7.text = "PHÂN TÍCH QUY TẮC PHẦN CỨNG TẠI SLIDE 6 BÀI GIẢNG GỐC:"
    p7.font.name = "Arial"; p7.font.size = Pt(13); p7.font.bold = True; p7.font.color.rgb = COLOR_DARK
    
    p7_items = [
        ("Đồng thời bên trong mỗi nhân", "Mỗi nhân CPU vẫn liên tục chuyển đổi xen kẽ giữa các luồng để đảm bảo tiến triển liên tục."),
        ("Không trùng lặp task trên 2 nhân", "Hệ điều hành đảm bảo tại 1 thời điểm không gán cùng 1 tập lệnh của 1 task cho 2 nhân khác nhau chạy trùng."),
        ("Tổng thể hệ thống", "Sự kết hợp hoàn hảo: Song song ở tầng vĩ mô (đa nhân) và Đồng thời ở tầng vi mô (nội bộ nhân).")
    ]
    for k, v in p7_items:
        p = tf_7.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run(); r1.text = f"• {k}: "; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = COLOR_GREEN
        r2 = p.add_run(); r2.text = v; r2.font.size = Pt(11); r2.font.color.rgb = COLOR_BODY

    add_speaker_notes(s7, "Slide 6 bài giảng gốc lưu ý điểm rất tinh tế: Trong thực tế, trên mỗi nhân CPU vẫn diễn ra xử lý đồng thời để chạy các tiến trình nền của OS, miễn là không chạy trùng 1 task trên 2 nhân cùng lúc.")

    # =========================================================================
    # SLIDE 8: XỬ LÝ BẤT ĐỒNG BỘ (ASYNCHRONOUS) — KHÁI NIỆM & ĐẶC ĐIỂM (Trang 7 slide gốc)
    # =========================================================================
    s8 = create_slide_base(prs, 8, "GIỚI THIỆU", "Asynchronous - Bất đồng bộ (Slide 7 gốc)")
    add_content_title(s8, "Bất đồng bộ (Asynchronous Computing)")
    add_bullets(s8, [
        ("Xử lý không theo thứ tự", "Cho phép các task được xử lý KHÔNG THEO THỨ TỰ (có thể chuyển sang task khác trước khi task trước đó hoàn thành)."),
        ("Giải quyết nhiều yêu cầu đồng thời", "Hệ thống có thể tiếp nhận và xử lý nhiều yêu cầu cùng lúc mà không bị nghẽn lại ở bất kỳ yêu cầu đơn lẻ nào."),
        ("Rút ngắn thời gian", "Hoàn thành toàn bộ khối lượng công việc trong khoảng thời gian ngắn hơn rất nhiều so với tuần tự."),
        ("Ví dụ đời sống trong slide gốc", "VD: Đi thi (làm bài thi), Bãi giữ xe (quét thẻ gửi xe), ...")
    ], top=Pt(150.0), height=Pt(165.0), font_size=Pt(12.5))
    s8.shapes.add_picture(os.path.join(IMG_DIR, "diag_orig_async_parking_exam.png"), Pt(28.34), Pt(325.0), Pt(735.0), Pt(195.0))

    add_speaker_notes(s8, "Slide 7 bài giảng gốc giới thiệu Bất đồng bộ: Cho phép các task xử lý không theo thứ tự, chuyển sang task khác trước khi task cũ xong. Giảng viên đưa ra 2 ví dụ đời sống rất hay: Đi thi và Bãi giữ xe.")

    # =========================================================================
    # SLIDE 9: PHÂN TÍCH HAI VÍ DỤ ĐỜI SỐNG: ĐI THI & BÃI GIỮ XE
    # =========================================================================
    s9 = create_slide_base(prs, 9, "VÍ DỤ THỰC TẾ", "Hai Ví dụ Đời sống trong Bài giảng gốc C4")
    add_content_title(s9, "Làm rõ bản chất Bất đồng bộ qua 2 ví dụ thực tế")
    add_bullets(s9, [
        ("Ví dụ 1 — Đi thi (Làm bài thi)", "Gặp câu hỏi khó hoặc cần thời gian suy nghĩ ➔ Thí sinh không ngồi chờ chết thời gian mà TẠM BỎ QUA chuyển sang làm câu dễ trước. Khi xong hoặc có ý tưởng mới quay lại làm câu khó. Đây là xử lý không theo thứ tự!"),
        ("Ví dụ 2 — Bãi giữ xe (Gửi xe & Lấy vé)", "Người gửi xe quét thẻ nhận vé (nhận handle) rồi tự lái xe vào bãi tìm chỗ đỗ. Nhân viên bảo vệ không cần đi theo xe mà TIẾP TỤC QUÉT THẺ phục vụ người tiếp theo. Nhiều yêu cầu được giải quyết đồng thời mà không ai bị chặn!"),
        ("Bài học công nghệ", "Bất đồng bộ giúp chương trình máy tính không bị lãng phí thời gian chờ đợi các thao tác chậm chạp (như mạng hay đĩa).")
    ], top=Pt(150.0), height=Pt(170.0), font_size=Pt(12.0))

    t_ex_h = ["Đặc điểm so sánh", "Mô hình Tuần tự (Synchronous Blocking)", "Mô hình Bất đồng bộ (Asynchronous Non-blocking)"]
    t_ex_r = [
        ["Trong phòng thi", "Ngồi chờ giải xong câu 1 mới được đọc câu 2 ➔ Hết giờ thi vẫn chưa làm xong!", "Làm câu dễ trước, câu khó làm sau ➔ Tối ưu hóa điểm số và thời gian làm bài!"],
        ["Tại bãi giữ xe", "Bảo vệ đi theo từng xe cho đến khi đỗ xong mới nhận xe tiếp ➔ Ùn tắc kéo dài!", "Bảo vệ phát vé rồi nhận xe tiếp theo ngay ➔ Bãi xe thông suốt, phục vụ hàng nghìn người!"]
    ]
    add_table_box(s9, t_ex_h, t_ex_r, Pt(28.34), Pt(335.0), Pt(735.0), Pt(180.0), col_widths=[160, 285, 290])

    add_speaker_notes(s9, "Hai ví dụ đời sống trong slide gốc giúp sinh viên hiểu ngay bản chất: Bất đồng bộ là không chờ đợi một cách vô ích. Việc nào làm được trước thì làm trước, việc nào phải chờ thì nhận vé hẹn làm sau.")

    # =========================================================================
    # SLIDE 10: MÔ HÌNH BẤT ĐỒNG BỘ THỰC THI ĐỒNG THỜI (Trang 8 slide gốc)
    # =========================================================================
    s10 = create_slide_base(prs, 10, "GIỚI THIỆU", "Mô hình Bất đồng bộ Thực thi Đồng thời (Slide 8 gốc)")
    add_content_title(s10, "Mô hình bất đồng bộ thực thi đồng thời trên một luồng điều khiển")
    add_bullets(s10, [
        ("Thực hiện xen kẽ trong một khoảng thời gian", "Trong một khoảng thời gian xác định, các task được thực hiện xen kẽ nhau linh hoạt."),
        ("Nằm trong một luồng điều khiển duy nhất", "Toàn bộ mô hình chạy trong MỘT LUỒNG ĐIỀU KHIỂN CHÍNH DUY NHẤT (Single Thread of Control)."),
        ("Nguyên tắc thực thi luồng đơn", "Khi 1 task đang thực thi thì các task khác KHÔNG THỰC THI; việc thực thi của 1 task có thể bị TẠM DỪNG và sau đó TIẾP TỤC LẠI."),
        ("Phạm vi ứng dụng rộng rãi", "Áp dụng hiệu quả cả trong hệ thống đơn bộ xử lý (Uniprocessor systems) lẫn hệ thống đa bộ xử lý (Multiprocessor systems).")
    ], top=Pt(150.0), height=Pt(165.0), font_size=Pt(12.5))
    s10.shapes.add_picture(os.path.join(IMG_DIR, "diag_orig_async_single_thread.png"), Pt(28.34), Pt(325.0), Pt(735.0), Pt(195.0))

    add_speaker_notes(s10, "Slide 8 bài giảng gốc là slide kết luận cực kỳ quan trọng của phần Giới thiệu: Mô hình bất đồng bộ thực thi đồng thời nằm trong MỘT LUỒNG DUY NHẤT. Task có thể tạm dừng rồi tiếp tục lại sau, áp dụng được cho cả máy 1 CPU lẫn đa CPU.")

    # =========================================================================
    # SLIDE 11: TỔNG KẾT PHẦN GIỚI THIỆU & CẦU NỐI CHUYỂN GIAO SANG PHẦN II
    # =========================================================================
    s11 = create_slide_base(prs, 11, "TỔNG KẾT PHẦN I", "Tổng kết Phần Giới thiệu & Bàn giao Kỹ thuật")
    add_content_title(s11, "Tóm tắt nền tảng & Chuyển giao sang Module concurrent.futures")
    add_bullets(s11, [
        ("Hoàn thành mục tiêu Phần I (TV1)", "Đã phân định chính xác và rõ ràng 4 mô hình: Tuần tự (theo thứ tự), Đồng thời (xen kẽ trên 1 core), Song song (cùng lúc trên đa core), và Bất đồng bộ (không theo thứ tự trên 1 luồng đơn)."),
        ("Quy luật cốt lõi cần nhớ", "Bất đồng bộ tận dụng thời gian rảnh khi một tác vụ tạm dừng để phục vụ tác vụ khác, giúp hoàn thành công việc nhanh hơn mà không nhất thiết phải dùng nhiều luồng."),
        ("Cầu nối chuyển giao sang Phần II", "Từ nguyên lý này, Python cung cấp công cụ gì để triển khai trong thực tế? Đó chính là Module `concurrent.futures`!"),
        ("Chuyển giao cho Thành viên 2 & 3", "Xin kính mời bạn Thành viên 2 tiếp nối bài giảng với chuyên đề: 'Sử dụng Module concurrent.futures (ThreadPoolExecutor & ProcessPoolExecutor)'!")
    ], top=Pt(150.0), height=Pt(180.0), font_size=Pt(12.0))

    c_box11 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(28.34), Pt(350.0), Pt(735.0), Pt(165.0))
    c_box11.fill.solid()
    c_box11.fill.fore_color.rgb = COLOR_CARD_BG
    c_box11.line.color.rgb = COLOR_RED
    c_box11.line.width = Pt(1.5)
    tf_11 = c_box11.text_frame
    tf_11.margin_left = Pt(20); tf_11.margin_top = Pt(15)
    p11 = tf_11.paragraphs[0]
    p11.text = "MẠCH NỐI CHUYỂN TIẾP TRONG FILE BÀI GIẢNG GỐC C4:"
    p11.font.name = "Arial"; p11.font.size = Pt(13); p11.font.bold = True; p11.font.color.rgb = COLOR_RED
    
    p11_items = [
        ("Phần I (Slide 1 đến 8 - TV1 vừa trình bày)", "Nền tảng lý thuyết: Tuần tự ➔ Đồng thời ➔ Song song ➔ Bất đồng bộ."),
        ("Phần II (Slide 9 đến 15 - TV2 & TV3)", "Sử dụng Module concurrent.futures: Executor, Future, ThreadPool và ProcessPool."),
        ("Phần III (Slide 16 đến 22 - TV4 & TV5)", "Quản lý vòng lặp sự kiện và xử lý tác vụ đồng thời với Asyncio."),
        ("Phần IV (Slide 23 đến 24 - TV6)", "Lưu ý thực tế và phối hợp các mô hình tính toán.")
    ]
    for k, v in p11_items:
        p = tf_11.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run(); r1.text = f"• {k}: "; r1.font.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = COLOR_DARK
        r2 = p.add_run(); r2.text = v; r2.font.size = Pt(10.5); r2.font.color.rgb = COLOR_BODY

    add_speaker_notes(s11, "Em xin kết thúc phần trình bày của Thành viên 1. Toàn bộ phần Giới thiệu từ slide 1 đến slide 8 đã được làm rõ. Xin mời bạn Thành viên 2 tiếp tục với module concurrent.futures ở slide 9!")

    # Lưu file
    out_file = "TV1_BaiGiang_ChuanFileGoc_C4.pptx"
    prs.save(out_file)
    print(f"XUẤT BẢN THÀNH CÔNG BỘ SLIDE CHUẨN FILE GỐC: {out_file} (Kích thước: {os.path.getsize(out_file):,} bytes)")

    try:
        prs.save("TV1_BaiGiang_TinhToanSongSong_PhanTan_C4.pptx")
        print("ĐÃ CẬP NHẬT ĐÈ VÀO: TV1_BaiGiang_TinhToanSongSong_PhanTan_C4.pptx")
    except PermissionError:
        print("LƯU Ý: TV1_BaiGiang_TinhToanSongSong_PhanTan_C4.pptx đang mở trong PowerPoint, bản chuẩn file gốc lưu tại TV1_BaiGiang_ChuanFileGoc_C4.pptx")

def main():
    prs = Presentation()
    prs.slide_width = Pt(793.8)
    prs.slide_height = Pt(595.2)
    build_tv1_clean_presentation(prs)

if __name__ == "__main__":
    main()
